import docx
from docx import Document

def fix_documents():
    file_path = "CLCO_WEEK1.docx"
    doc = Document(file_path)

    # 1. Sửa trong paragraphs: thay thế 25110159 thành 24110354 và /dev/ thành /main/
    for p in doc.paragraphs:
        if "25110159" in p.text:
            for r in p.runs:
                if "25110159" in r.text:
                    r.text = r.text.replace("25110159", "24110354")
        if "/blob/dev/" in p.text or "/tree/dev/" in p.text:
            for r in p.runs:
                r.text = r.text.replace("/blob/dev/", "/blob/main/").replace("/tree/dev/", "/tree/main/")

    # 2. Sửa trong tất cả các bảng
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if "25110159" in cell.text:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            if "25110159" in r.text:
                                r.text = r.text.replace("25110159", "24110354")
                if "/dev/" in cell.text or "Nhánh dev" in cell.text:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.text = (
                                r.text.replace("/blob/dev/", "/blob/main/")
                                .replace("/tree/dev/", "/tree/main/")
                                .replace("Nhánh dev", "Nhánh main")
                                .replace("nhánh dev", "nhánh main")
                            )

    doc.save("CLCO_WEEK1.docx")
    doc.save("CLCO_WEEK1_fix.docx")
    print("Successfully updated CLCO_WEEK1.docx and created clean CLCO_WEEK1_fix.docx!")

if __name__ == "__main__":
    fix_documents()
