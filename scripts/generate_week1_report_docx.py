import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = Document()
    
    # Page Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Style configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(30, 30, 30)

    # --- HEADER / TITLE ---
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_inst.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — TRƯỜNG ĐẠI HỌC SƯ PHẠM KỸ THUẬT TP.HCM\nKHOA CÔNG NGHỆ THÔNG TIN\n---------------------------------------")
    run_inst.font.size = Pt(11)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(80, 80, 80)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(16)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run("BÁO CÁO TIẾN ĐỘ THỰC HIỆN TUẦN 1")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 74, 150)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("ĐỀ TÀI: HỆ THỐNG QUẢN LÝ BÃI GIỮ XE THÔNG MINH\n(CLOUD-BASED SMART PARKING SYSTEM)")
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(40, 40, 40)

    # Info Box Table
    info_table = doc.add_table(rows=6, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False

    data_info = [
        ("Học phần:", "Điện toán đám mây (Cloud Computing) — Đợt 1 (2026 - 2027)"),
        ("Giảng viên hướng dẫn:", "ThS./TS. Huỳnh Xuân Phụng"),
        ("Sinh viên thực hiện 1:", "Lê Cao Việt Tín (Trưởng nhóm - Kiến trúc Serverless, AI & Backend Core)"),
        ("Sinh viên thực hiện 2:", "Nguyễn Trung Hải — MSSV: 24110207 (Hạ tầng AWS, Dataset & QA)"),
        ("Môi trường triển khai:", "AWS Learner Lab | Region: us-east-1 (N. Virginia)"),
        ("GitHub Repository:", "https://github.com/nthai212006-gh/cloud-smart-parking-system (Nhánh dev / main)")
    ]

    for i, (k, v) in enumerate(data_info):
        row = info_table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = RGBColor(20, 50, 100)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(11)
        if "github.com" in v:
            r1.bold = True
            r1.font.color.rgb = RGBColor(0, 102, 204)
        
        set_cell_background(c0, "F0F4F8")
        set_cell_background(c1, "FAFCFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Function to add heading 1
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(16, 74, 150)
        return h

    # Function to add heading 2
    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(40, 90, 140)
        return h

    # --- 1. TỔNG QUAN ĐỀ TÀI & MỤC TIÊU TUẦN 1 ---
    add_h1("1. TỔNG QUAN ĐỀ TÀI VÀ MỤC TIÊU TUẦN 1")
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.add_run(
        "Đề tài hướng tới việc xây dựng hệ thống quản lý bãi giữ xe thông minh ứng dụng kiến trúc không máy chủ (Serverless Architecture) "
        "trên nền tảng điện toán đám mây Amazon Web Services (AWS). Hệ thống tự động nhận dạng biển số xe (ô tô, xe máy) bằng thị giác máy tính "
        "thông qua AWS Rekognition, điều phối dữ liệu lượt xe vào/ra tức thời bằng AWS Lambda và API Gateway, lưu trữ dữ liệu an toàn trên Amazon DynamoDB và S3, "
        "đồng thời phân tích thống kê báo cáo qua Amazon Athena.\n\n"
        "Trong Tuần 1 (06/10/2026 - 12/10/2026), nhóm tập trung hoàn thành các mục tiêu nền móng quan trọng:\n"
        "• Nghiên cứu giải pháp và thiết kế kiến trúc Serverless toàn diện (Sơ đồ kiến trúc, Data Flow, Database Schema).\n"
        "• Thiết kế và hiện thực hóa đặc tả API Gateway (OpenAPI/Swagger) cùng 2 hàm AWS Lambda cốt lõi: SmartParking-Checkin và SmartParking-Checkout.\n"
        "• Xây dựng và kiểm thử cục bộ giải thuật bóc tách, chuẩn hóa biển số xe Việt Nam (test_plate_regex.py).\n"
        "• Thiết lập hạ tầng AWS Learner Lab: Tạo S3 Bucket, cấu hình cấu trúc thư mục, khởi tạo bảng DynamoDB ParkingTickets, tuân thủ nghiêm ngặt chuẩn Tagging (Project, Owner, Environment).\n"
        "• Thu thập, phân loại và tiền xử lý bộ dữ liệu 120 ảnh biển số xe thử nghiệm, tuyệt đối tuân thủ quy chuẩn bảo mật PII (che mờ khuôn mặt người và thông tin nhạy cảm).\n"
        "• Quản trị mã nguồn qua Git: Phân nhánh dev, pull request rà soát mã nguồn minh bạch."
    )

    # --- 2. LIÊN KẾT MÃ NGUỒN GITHUB & CẤU TRÚC REPOSITORY ---
    add_h1("2. LIÊN KẾT MÃ NGUỒN GITHUB VÀ CẤU TRÚC THƯ MỤC")
    p_git = doc.add_paragraph()
    p_git.paragraph_format.line_spacing = 1.2
    p_git.add_run("• Đường dẫn Repository GitHub: ").bold = True
    r_git_link = p_git.add_run("https://github.com/nthai212006-gh/cloud-smart-parking-system\n")
    r_git_link.font.color.rgb = RGBColor(0, 102, 204)
    r_git_link.bold = True
    p_git.add_run("• Trạng thái truy cập: ").bold = True
    p_git.add_run("Công khai (Public) — Giảng viên và người đánh giá có thể truy cập, xem code và clone trực tiếp.\n")
    p_git.add_run("• Nhánh hoạt động: ").bold = True
    p_git.add_run("dev (nhánh phát triển chính với đầy đủ commit history) và main (nhánh đồng bộ nghiệm thu).\n\n")
    p_git.add_run("Cấu trúc cây thư mục mã nguồn trên Git:\n").bold = True

    p_tree = doc.add_paragraph()
    p_tree.paragraph_format.left_indent = Inches(0.3)
    run_tree = p_tree.add_run(
        "cloud-smart-parking-system/\n"
        "├── backend/\n"
        "│   ├── api_gateway_openapi.json    # Đặc tả REST API Gateway (OpenAPI 3.0 / Swagger)\n"
        "│   ├── checkin.py                  # Lambda xử lý xe vào (SmartParking-Checkin)\n"
        "│   ├── checkout.py                 # Lambda xử lý xe ra & tính tiền (SmartParking-Checkout)\n"
        "│   └── README.md                   # Mô tả chi tiết giao tiếp tầng Backend\n"
        "├── docs/\n"
        "│   ├── THIET_KE_KIEN_TRUC_HE_THONG.md  # Tài liệu kiến trúc hệ thống Serverless\n"
        "│   ├── architecture/               # Bản vẽ kiến trúc & cấu trúc bảng dữ liệu\n"
        "│   ├── guidelines/                 # Quy chuẩn Tagging & đặt tên tài nguyên AWS\n"
        "│   ├── plans/week-01/              # Bản kế hoạch phân công chi tiết Tuần 1\n"
        "│   └── worklogs/week-01/           # Nhật ký thực hiện công việc Tuần 1\n"
        "├── evidence/week1/                 # Ảnh chụp màn hình minh chứng cấu hình AWS thật\n"
        "│   ├── s3_console.png, s3_folders.png, dynamodb_console.png, billing_cost.png\n"
        "├── scripts/\n"
        "│   ├── test_plate_regex.py         # Kiểm thử thuật toán Regex OCR biển số xe\n"
        "│   ├── setup_aws_resources.py      # Script tự động hóa khởi tạo S3 & DynamoDB\n"
        "│   └── build_curated_dataset.py    # Xử lý và phân loại tập dữ liệu ảnh\n"
        "├── test-images/                    # Bộ 120 ảnh thử nghiệm đã che mờ khuôn mặt (PII Privacy)\n"
        "└── README.md                       # Tài liệu tổng quan dự án"
    )
    run_tree.font.name = "Consolas"
    run_tree.font.size = Pt(9.5)

    # --- 3. THIẾT KẾ KIẾN TRÚC HỆ THỐNG SERVERLESS ---
    add_h1("3. THIẾT KẾ KIẾN TRÚC HỆ THỐNG SERVERLESS")
    p_arch = doc.add_paragraph()
    p_arch.paragraph_format.line_spacing = 1.2
    p_arch.add_run(
        "Hệ thống vận hành theo mô hình kiến trúc hướng sự kiện (Event-driven Serverless Architecture) trên AWS. "
        "Mô hình này giúp tối ưu hóa chi phí (Pay-as-you-go), không mất chi phí duy trì máy chủ khi bãi xe không có lượt ra vào, "
        "và tự động mở rộng (auto-scaling) tức thời khi lưu lượng xe tăng đột biến.\n\n"
    )

    # Table of Services
    svc_table = doc.add_table(rows=6, cols=3)
    svc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    svc_table.autofit = False

    svc_headers = ["Thành phần / Dịch vụ AWS", "Tên tài nguyên", "Vai trò và Chức năng kỹ thuật"]
    for j, h in enumerate(svc_headers):
        cell = svc_table.rows[0].cells[j]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "104A96")
        set_cell_margins(cell, 80, 80, 100, 100)

    svc_rows = [
        ("Amazon API Gateway", "SmartParking-REST-API", "Tiếp nhận request RESTful HTTPS (/checkin, /checkout), xác thực, cấu hình CORS, tích hợp proxy chuyển tiếp sự kiện đến AWS Lambda."),
        ("AWS Lambda", "SmartParking-Checkin\nSmartParking-Checkout", "Tầng tính toán serverless (Python 3.11): Thực thi logic nghiệp vụ xe vào/ra, điều phối gọi Rekognition, cập nhật DynamoDB và lưu trữ S3."),
        ("Amazon Rekognition", "DetectText API", "Dịch vụ thị giác AI nhận dạng ký tự (OCR). Bóc tách văn bản trên biển số xe máy (2 dòng) và ô tô (1 dòng) với chỉ số tin cậy Confidence > 80%."),
        ("Amazon DynamoDB", "Table: ParkingTickets", "Cơ sở dữ liệu NoSQL lưu trữ bản ghi vé xe theo thời gian thực (Partition Key: ticketId). Billing mode On-Demand tối ưu chi phí."),
        ("Amazon S3", "smart-parking-images-24110207", "Kho lưu trữ đối tượng chứa ảnh xe vào (in/), ảnh xe ra (out/) và khu vực lưu trữ logs truy vấn của Athena (athena-logs/).")
    ]

    for i, row_data in enumerate(svc_rows):
        row = svc_table.rows[i+1]
        for j, val in enumerate(row_data):
            c = row.cells[j]
            c.width = Inches(1.8) if j == 0 else (Inches(1.8) if j == 1 else Inches(3.4))
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(10)
            if j == 0:
                r.bold = True
            bg = "F9FBFD" if i % 2 == 0 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 80, 80)

    # --- 4. THIẾT KẾ CƠ SỞ DỮ LIỆU & QUY CHUẨN LƯU TRỮ ---
    add_h1("4. THIẾT KẾ CƠ SỞ DỮ LIỆU VÀ QUY CHUẨN LƯU TRỮ")
    add_h2("4.1. Lược đồ bảng Amazon DynamoDB: ParkingTickets")
    p_db = doc.add_paragraph()
    p_db.paragraph_format.line_spacing = 1.2
    p_db.add_run(
        "Bảng DynamoDB được khởi tạo với chế độ On-Demand (PAY_PER_REQUEST) để đáp ứng chuẩn AWS Learner Lab, "
        "Partition Key là ticketId (kiểu String - định dạng UUID v4 đảm bảo tính duy nhất tuyệt đối).\n"
    )

    schema_table = doc.add_table(rows=9, cols=3)
    schema_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    schema_headers = ["Tên trường (Attribute)", "Kiểu dữ liệu", "Mô tả chi tiết và Quy chuẩn"]
    for j, h in enumerate(schema_headers):
        cell = schema_table.rows[0].cells[j]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "2C5E8A")
        set_cell_margins(cell, 80, 80, 100, 100)

    schema_rows = [
        ("ticketId (PK)", "String", "Mã vé xe duy nhất, sinh ngẫu nhiên theo chuẩn UUID v4."),
        ("plateNumber", "String", "Biển số xe đã qua chuẩn hóa regex (Ví dụ: 59X112345, 51A12345)."),
        ("vehicleType", "String", "Phân loại phương tiện gửi: MOTORBIKE (xe máy) hoặc CAR (ô tô)."),
        ("checkinTime", "String", "Thời điểm xe vào bãi định dạng chuẩn quốc tế ISO 8601 (UTC, kết thúc bằng 'Z')."),
        ("checkoutTime", "String", "Thời điểm xe ra bãi (được cập nhật khi check-out hoàn tất)."),
        ("status", "String", "Trạng thái vé: PARKED (xe đang trong bãi) hoặc COMPLETED (đã lấy xe)."),
        ("fee", "Number", "Phí gửi xe tính bằng VNĐ. Lúc check-in mặc định = 0; tính tự động lúc check-out."),
        ("inImageKey / outImageKey", "String", "Đường dẫn đối tượng lưu trữ ảnh trong S3 (Ví dụ: in/ticket_123.jpg).")
    ]

    for i, row_data in enumerate(schema_rows):
        row = schema_table.rows[i+1]
        for j, val in enumerate(row_data):
            c = row.cells[j]
            c.width = Inches(1.8) if j == 0 else (Inches(1.2) if j == 1 else Inches(4.0))
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if j == 0:
                r.bold = True
                r.font.name = "Consolas"
            bg = "F9FBFD" if i % 2 == 0 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 50, 50, 80, 80)

    add_h2("4.2. Quy chuẩn lưu trữ trên Amazon S3")
    p_s3 = doc.add_paragraph()
    p_s3.paragraph_format.line_spacing = 1.2
    p_s3.add_run(
        "Bucket S3: smart-parking-images-24110207 (Region us-east-1). Kích hoạt Block All Public Access.\n"
        "Cấu trúc tiền tố phân vùng (Prefix Partitioning):\n"
        "• in/ : Lưu trữ ảnh chụp xe vào bãi tương ứng theo mã vé {ticketId}.jpg.\n"
        "• out/ : Lưu trữ ảnh chụp xe ra bãi phục vụ đối soát và xử lý tranh chấp.\n"
        "• athena-logs/ : Phân vùng dành riêng cho kết quả truy vấn nhật ký của Amazon Athena."
    )

    # --- 5. CHI TIẾT MÃ NGUỒN API GATEWAY & AWS LAMBDA ---
    add_h1("5. CHI TIẾT HIỆN THỰC API GATEWAY VÀ AWS LAMBDA")
    add_h2("5.1. Đặc tả Amazon API Gateway (OpenAPI 3.0)")
    p_gw = doc.add_paragraph()
    p_gw.paragraph_format.line_spacing = 1.2
    p_gw.add_run(
        "Nhóm đã xây dựng đặc tả chuẩn OpenAPI 3.0 (file backend/api_gateway_openapi.json) định nghĩa 2 REST API endpoints chính:\n"
        "• POST /checkin: Tiếp nhận payload JSON chứa ảnh Base64 và loại xe vehicleType. Tích hợp proxy chuyển tiếp tới Lambda SmartParking-Checkin. Thiết lập đầy đủ cấu hình phản hồi OPTIONS xử lý CORS.\n"
        "• POST /checkout: Tiếp nhận ticketId, kích hoạt Lambda SmartParking-Checkout đối soát thời gian và tính tiền. Hỗ trợ CORS đa nền tảng.\n"
    )

    add_h2("5.2. Hàm Lambda xử lý xe vào (backend/checkin.py)")
    p_code1 = doc.add_paragraph()
    p_code1.paragraph_format.line_spacing = 1.15
    p_code1.paragraph_format.left_indent = Inches(0.2)
    r_c1 = p_code1.add_run(
        "# Trích đoạn mã nguồn Lambda Check-in cốt lõi:\n"
        "def lambda_handler(event, context):\n"
        "    body = json.loads(event.get('body', '{}'))\n"
        "    image_base64 = body.get('image')\n"
        "    vehicle_type = body.get('vehicleType', 'MOTORBIKE')\n"
        "    \n"
        "    plate_number = extract_plate(image_base64)\n"
        "    if not plate_number:\n"
        "        plate_number = 'UNKNOWN'  # Gán cờ hỗ trợ nhập thủ công nếu ảnh mờ\n"
        "        \n"
        "    ticket_id = str(uuid.uuid4())\n"
        "    checkin_time = datetime.datetime.utcnow().isoformat() + 'Z'\n"
        "    # Ghi nhận vé vào DynamoDB ParkingTickets và lưu ảnh vào S3 in/\n"
        "    return {\n"
        "        'statusCode': 200,\n"
        "        'headers': {'Access-Control-Allow-Origin': '*', 'Content-Type': 'application/json'},\n"
        "        'body': json.dumps({'ticketId': ticket_id, 'plateNumber': plate_number, 'checkinTime': checkin_time, 'status': 'PARKED'})\n"
        "    }"
    )
    r_c1.font.name = "Consolas"
    r_c1.font.size = Pt(9.5)

    add_h2("5.3. Hàm Lambda xử lý xe ra & Thuật toán tính phí (backend/checkout.py)")
    p_code2 = doc.add_paragraph()
    p_code2.paragraph_format.line_spacing = 1.15
    p_code2.paragraph_format.left_indent = Inches(0.2)
    r_c2 = p_code2.add_run(
        "# Thuật toán tính tiền bãi xe theo block giờ thực tế:\n"
        "def calculate_fee(vehicle_type, checkin_time_str, checkout_time_str):\n"
        "    checkin = datetime.datetime.fromisoformat(checkin_time_str.replace('Z', '+00:00'))\n"
        "    checkout = datetime.datetime.fromisoformat(checkout_time_str.replace('Z', '+00:00'))\n"
        "    delta = checkout - checkin\n"
        "    total_hours = math.ceil(delta.total_seconds() / 3600)\n"
        "    if total_hours <= 0: total_hours = 1\n"
        "    \n"
        "    fee = 0\n"
        "    if vehicle_type == 'MOTORBIKE':\n"
        "        fee = 5000  # 4 giờ đầu 5.000 VNĐ\n"
        "        if total_hours > 4: fee += (total_hours - 4) * 2000  # Mỗi giờ tiếp theo +2.000 VNĐ\n"
        "    elif vehicle_type == 'CAR':\n"
        "        fee = 20000 # 2 giờ đầu 20.000 VNĐ\n"
        "        if total_hours > 2: fee += (total_hours - 2) * 10000 # Mỗi giờ tiếp theo +10.000 VNĐ\n"
        "    return fee"
    )
    r_c2.font.name = "Consolas"
    r_c2.font.size = Pt(9.5)

    add_h2("5.4. Giải thuật Regex bóc tách biển số xe (scripts/test_plate_regex.py)")
    p_regex = doc.add_paragraph()
    p_regex.paragraph_format.line_spacing = 1.2
    p_regex.add_run(
        "Nhóm đã xây dựng và kiểm thử thành công giải thuật trích xuất văn bản từ mảng TextDetections của Rekognition: "
        "Lọc các phần tử có Type == 'LINE' với Confidence > 80%, sau đó áp dụng biểu thức chính quy re.sub(r'[^A-Z0-9]', '', raw_text.upper()) "
        "để loại bỏ triệt để các ký tự gây nhiễu (dấu gạch nối, dấu chấm, khoảng trắng, bóng đổ camera), cho ra chuỗi biển số xe đồng nhất."
    )

    # --- 6. BỘ DỮ LIỆU THỬ NGHIỆM VÀ QUY CHUẨN BẢO MẬT PII ---
    add_h1("6. BỘ DỮ LIỆU THỬ NGHIỆM VÀ QUY CHUẨN BẢO MẬT PII")
    p_data = doc.add_paragraph()
    p_data.paragraph_format.line_spacing = 1.2
    p_data.add_run(
        "Nhóm đã thu thập và chọn lọc 120 ảnh thử nghiệm thực tế (thư mục test-images/), phân chia thành 5 nhóm kịch bản kiểm thử:\n"
        "1. 01_car_clear (40 ảnh): Biển số ô tô rõ nét, chụp ban ngày, góc chụp thẳng.\n"
        "2. 02_motor_clear (40 ảnh): Biển số xe máy rõ nét, cấu trúc 2 dòng tiêu chuẩn.\n"
        "3. 03_car_hard (10 ảnh): Biển số ô tô trong điều kiện thiếu sáng, góc nghiêng hoặc hơi mờ.\n"
        "4. 04_motor_hard (10 ảnh): Biển số xe máy bị dính bùn, mờ số hoặc chụp đêm.\n"
        "5. 05_error_cases (20 ảnh): Các tình huống ngoại lệ (không có biển số, ảnh quá tối, góc khuất) để kiểm tra khả năng phục hồi của Lambda.\n\n"
        "QUY CHUẨN BẢO MẬT (PII Privacy): 100% hình ảnh trong tập dữ liệu đều đã được xử lý làm mờ/bôi đen toàn bộ khuôn mặt người đi đường, "
        "người lái xe và các thông tin nhạy cảm trước khi đưa vào hệ thống và đưa lên Git."
    )

    # --- 7. BẢNG PHÂN CÔNG CÔNG VIỆC VÀ TỔNG HỢP GIỜ LÀM VIỆC ---
    add_h1("7. BẢNG PHÂN CÔNG CÔNG VIỆC VÀ TỔNG HỢP GIỜ LÀM VIỆC")
    p_hours_note = doc.add_paragraph()
    p_hours_note.paragraph_format.line_spacing = 1.2
    p_hours_note.add_run(
        "Theo quy định môn học Điện toán đám mây, mỗi thành viên phải đảm bảo tối thiểu 8.0 giờ làm việc thực tế trong tuần. "
        "Cả hai thành viên trong nhóm đều đã hoàn thành và vượt chỉ tiêu đề ra:\n"
    )

    hr_table = doc.add_table(rows=4, cols=4)
    hr_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hr_table.autofit = False

    hr_headers = ["Thành viên thực hiện", "Nhiệm vụ chính phụ trách", "Tổng giờ", "Đánh giá"]
    for j, h in enumerate(hr_headers):
        cell = hr_table.rows[0].cells[j]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "104A96")
        set_cell_margins(cell, 80, 80, 100, 100)

    hr_rows = [
        ("Lê Cao Việt Tín\n(Trưởng nhóm)", "Nghiên cứu AWS Rekognition; Viết thuật toán Regex bóc tách biển số; Thiết kế kiến trúc Serverless; Viết code Lambda checkin/checkout; Cấu hình API Gateway OpenAPI spec.", "8.5 giờ\n(Quy định: ≥ 8h)", "ĐẠT\n(Vượt chỉ tiêu)"),
        ("Nguyễn Trung Hải\n(MSSV: 24110207)", "Khởi tạo hạ tầng S3 & DynamoDB trên Learner Lab; Thiết lập chuẩn Tagging; Thu thập và xử lý 120 ảnh test (che mờ PII mặt người); Viết script tự động hóa; Chụp minh chứng console.", "8.5 giờ\n(Quy định: ≥ 8h)", "ĐẠT\n(Vượt chỉ tiêu)"),
        ("TỔNG CỘNG CẢ NHÓM", "Hoàn thành 100% mục tiêu Tuần 1, mã nguồn sẵn sàng trên Git dev/main.", "17.0 giờ", "HOÀN THÀNH TỐT")
    ]

    for i, row_data in enumerate(hr_rows):
        row = hr_table.rows[i+1]
        for j, val in enumerate(row_data):
            c = row.cells[j]
            c.width = Inches(1.8) if j == 0 else (Inches(3.2) if j == 1 else (Inches(1.0) if j == 2 else Inches(1.0)))
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if i == 2 or j == 0:
                r.bold = True
            bg = "EDF4FA" if i == 2 else ("F9FBFD" if i % 2 == 0 else "FFFFFF")
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 80, 80)

    # --- 8. MINH CHỨNG THỰC TẾ TRÊN AWS CONSOLE ---
    add_h1("8. HÌNH ẢNH MINH CHỨNG CẤU HÌNH TRÊN AWS CONSOLE")
    p_ev_note = doc.add_paragraph()
    p_ev_note.paragraph_format.line_spacing = 1.2
    p_ev_note.add_run(
        "Dưới đây là các ảnh chụp màn hình thực tế trên AWS Management Console của nhóm trong môi trường AWS Learner Lab, "
        "thể hiện rõ Account ID, Region us-east-1, đồng hồ hệ thống và các tài nguyên đã khởi tạo thành công:"
    )

    ev_dir = r"c:\Users\T14 GEN2\Documents\cloud-smart-parking-system\evidence\week1"
    ev_images = [
        ("s3_console.png", "Hình 1: Amazon S3 Bucket 'smart-parking-images-24110207' tại Region us-east-1 với Block All Public Access và chuẩn Tagging."),
        ("s3_folders.png", "Hình 2: Cấu trúc các thư mục phân vùng in/, out/, athena-logs/ bên trong S3 Bucket."),
        ("dynamodb_console.png", "Hình 3: Bảng Amazon DynamoDB 'ParkingTickets' (Partition Key ticketId, Billing mode On-Demand, trạng thái Active)."),
        ("billing_cost.png", "Hình 4: Quản lý chi phí AWS Billing ($0.00 / $100.00 credits trong Learner Lab, an toàn tuyệt đối).")
    ]

    for img_name, caption in ev_images:
        img_path = os.path.join(ev_dir, img_name)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(2)
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(6.2))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(12)
            r_cap = p_cap.add_run(caption)
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(80, 80, 80)

    # --- 9. KHAI BÁO MINH BẠCH SỬ DỤNG TRÍ TUỆ NHÂN TẠO (AI TRANSPARENCY) ---
    add_h1("9. KHAI BÁO SỬ DỤNG TRÍ TUỆ NHÂN TẠO (AI TRANSPARENCY)")
    p_ai = doc.add_paragraph()
    p_ai.paragraph_format.line_spacing = 1.2
    p_ai.add_run(
        "Thực hiện theo đúng quy định học phần: Sinh viên được phép sử dụng công cụ AI hỗ trợ nhưng phải khai báo minh bạch "
        "và chịu trách nhiệm làm chủ toàn bộ nội dung, sẵn sàng giải thích chi tiết trong buổi vấn đáp của Giảng viên hướng dẫn:\n\n"
        "• Các công cụ AI đã tham khảo: Google Gemini, Anthropic Claude, ChatGPT.\n"
        "• Phạm vi sử dụng AI hỗ trợ:\n"
        "  1. Gợi ý cú pháp Mermaid cho sơ đồ luồng dữ liệu (sequenceDiagram, flowchart) trong tài liệu kỹ thuật.\n"
        "  2. Tham khảo các mẫu regex chuẩn hóa chuỗi bóc tách biển số từ kết quả thô của AWS Rekognition.\n"
        "  3. Gợi ý cấu trúc khung chuẩn OpenAPI 3.0 / Swagger cho API Gateway.\n"
        "• Mức độ làm chủ của Sinh viên:\n"
        "  100% tài nguyên AWS (S3, DynamoDB), code hàm Lambda và bộ dữ liệu 120 ảnh đã được hai thành viên trực tiếp thiết lập, "
        "chạy thử nghiệm và kiểm soát hoàn toàn trên môi trường AWS Learner Lab thực tế. Cả hai sinh viên hiểu sâu kiến trúc và cam kết tự tin trả lời vấn đáp."
    )

    # --- 10. KẾ HOẠCH TUẦN 2 ---
    add_h1("10. KẾ HOẠCH TRIỂN KHAI TUẦN 2")
    p_w2 = doc.add_paragraph()
    p_w2.paragraph_format.line_spacing = 1.2
    p_w2.add_run(
        "Dựa trên nền tảng Tuần 1 đã nghiệm thu, kế hoạch Tuần 2 của nhóm bao gồm:\n"
        "• Triển khai chính thức 2 hàm Lambda lên AWS Lambda Console kết nối IAM LabRole.\n"
        "• Deploy REST API Gateway lên Stage 'dev' và cấp quyền gọi Lambda (Lambda permissions).\n"
        "• Tích hợp trực tiếp Amazon Rekognition DetectText với ảnh thực tế trên S3.\n"
        "• Viết test suite Postman / Python script kiểm thử tự động toàn bộ luồng Check-in và Check-out với 120 ảnh mẫu."
    )

    # Signatures
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(20)
    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0 = sig_table.rows[0].cells[0]
    c1 = sig_table.rows[0].cells[1]
    c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s0 = c0.paragraphs[0].add_run("XÁC NHẬN CỦA SINH VIÊN 1\n(Trưởng nhóm)\n\n\n\nLê Cao Việt Tín")
    r_s1 = c1.paragraphs[0].add_run("XÁC NHẬN CỦA SINH VIÊN 2\n\n\n\n\nNguyễn Trung Hải")
    r_s0.bold = True
    r_s1.bold = True

    # Save
    out_path = r"c:\Users\T14 GEN2\Documents\cloud-smart-parking-system\BAO_CAO_TUAN_1_SMART_PARKING.docx"
    doc.save(out_path)
    print(f"Report generated successfully at: {out_path}")

if __name__ == "__main__":
    create_report()
